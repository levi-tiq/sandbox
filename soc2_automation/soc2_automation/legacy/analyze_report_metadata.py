from docx import Document
import re
from datetime import datetime
import os

def analyze_report_metadata(docx_path):
    """
    Analyze a SOC2 report to find company name and completion date patterns
    """
    print(f"🔍 Analyzing: {os.path.basename(docx_path)}")
    print("="*60)
    
    doc = Document(docx_path)
    
    # Look for patterns in the first few paragraphs
    print("\n📝 First 20 paragraphs:")
    for i, para in enumerate(doc.paragraphs[:20]):
        text = para.text.strip()
        if text:
            print(f"  {i+1:2d}: {text}")
    
    # Look for tables (often contain metadata)
    print(f"\n📊 First few tables:")
    for table_idx, table in enumerate(doc.tables[:3]):
        print(f"\nTable {table_idx + 1} ({len(table.rows)} rows, {len(table.columns)} cols):")
        
        for row_idx, row in enumerate(table.rows[:5]):  # First 5 rows
            row_text = " | ".join([cell.text.strip()[:80] for cell in row.cells])
            print(f"  Row {row_idx + 1}: {row_text}")
    
    # Look for date patterns throughout document
    print(f"\n📅 Date patterns found:")
    date_patterns = [
        r'\b\d{1,2}[/-]\d{1,2}[/-]\d{4}\b',  # MM/DD/YYYY or MM-DD-YYYY
        r'\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b',  # YYYY/MM/DD or YYYY-MM-DD
        r'\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{1,2},?\s+\d{4}\b',  # Month DD, YYYY
        r'\b\d{1,2}\s+(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}\b',  # DD Month YYYY
    ]
    
    found_dates = set()
    for para in doc.paragraphs:
        text = para.text
        for pattern in date_patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            for match in matches:
                found_dates.add(match)
    
    for date in sorted(found_dates):
        print(f"  📅 {date}")
    
    # Look for company name patterns
    print(f"\n🏢 Potential company names (from filename and content):")
    
    # Extract from filename
    filename = os.path.basename(docx_path)
    filename_company = filename.split('-')[0] if '-' in filename else filename.split('_')[0]
    print(f"  From filename: {filename_company}")
    
    # Look for patterns in first few paragraphs that might be company names
    company_keywords = ['company', 'corporation', 'corp', 'inc', 'llc', 'ltd', 'technologies', 'systems', 'services']
    
    for i, para in enumerate(doc.paragraphs[:10]):
        text = para.text.strip()
        if any(keyword.lower() in text.lower() for keyword in company_keywords):
            print(f"  Para {i+1}: {text}")

if __name__ == "__main__":
    # Analyze a sample report
    sample_report = "/home/levi/sandbox/soc2_automation/report_analysis/reports/Colligo_SOC2T2_Report_QA_merged_final.docx"
    
    if os.path.exists(sample_report):
        analyze_report_metadata(sample_report)
    else:
        print(f"Sample report not found: {sample_report}")
        
        # Try to find any report
        reports_dir = "/home/levi/sandbox/soc2_automation/report_analysis/reports/"
        if os.path.exists(reports_dir):
            reports = [f for f in os.listdir(reports_dir) if f.endswith('.docx')]
            if reports:
                sample_report = os.path.join(reports_dir, reports[0])
                print(f"Using first available report: {sample_report}")
                analyze_report_metadata(sample_report)