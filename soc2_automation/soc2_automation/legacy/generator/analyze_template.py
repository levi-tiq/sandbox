from docx import Document
from docx.oxml.ns import qn
import zipfile
import os

def analyze_template_images(docx_path):
    """
    Analyze the DOCX template to find all images and their locations
    """
    print(f"Analyzing template: {docx_path}")
    doc = Document(docx_path)
    
    def analyze_container(container, container_name):
        """Analyze a container (document, header, footer) for images"""
        image_count = 0
        
        for i, paragraph in enumerate(container.paragraphs):
            for j, run in enumerate(paragraph.runs):
                # Check for embedded images (simplified approach)
                drawing_elements = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                
                if drawing_elements:
                    image_count += 1
                    print(f"  📸 Image found in {container_name} - Paragraph {i+1}, Run {j+1}")
                    
                    # Try to get image dimensions if possible
                    drawings = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
                    
                    if drawings:
                        for drawing in drawings:
                            cx = drawing.get('cx')  # width in EMUs
                            cy = drawing.get('cy')  # height in EMUs
                            if cx and cy:
                                # Convert EMUs to inches (1 inch = 914400 EMUs)
                                width_inches = int(cx) / 914400
                                height_inches = int(cy) / 914400
                                print(f"    Size: {width_inches:.2f}\" x {height_inches:.2f}\"")
        
        return image_count
    
    total_images = 0
    
    # Analyze main document
    print("\n🔍 Main Document:")
    total_images += analyze_container(doc, "Main Document")
    
    # Analyze headers and footers
    for i, section in enumerate(doc.sections):
        print(f"\n🔍 Section {i+1} Header:")
        total_images += analyze_container(section.header, f"Section {i+1} Header")
        
        print(f"\n🔍 Section {i+1} Footer:")
        total_images += analyze_container(section.footer, f"Section {i+1} Footer")
    
    print(f"\n📊 Total images found: {total_images}")
    
    # Also check what images are stored in the DOCX file
    print("\n📁 Images stored in DOCX:")
    try:
        with zipfile.ZipFile(docx_path, 'r') as docx_zip:
            image_files = [f for f in docx_zip.namelist() if f.startswith('word/media/')]
            for img_file in image_files:
                print(f"  📄 {img_file}")
    except Exception as e:
        print(f"  ❌ Could not read DOCX structure: {e}")
    
    return total_images

def get_image_positions(docx_path):
    """
    Get detailed positions of images for replacement targeting
    """
    doc = Document(docx_path)
    positions = []
    
    def scan_container(container, container_type, section_num=None):
        for para_idx, paragraph in enumerate(container.paragraphs):
            for run_idx, run in enumerate(paragraph.runs):
                drawing_elements = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                
                if drawing_elements:
                    # Get dimensions
                    width_emu = height_emu = None
                    drawings = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
                    
                    if drawings:
                        width_emu = drawings[0].get('cx')
                        height_emu = drawings[0].get('cy')
                    
                    position = {
                        'container_type': container_type,
                        'section_num': section_num,
                        'paragraph_index': para_idx,
                        'run_index': run_idx,
                        'width_emu': width_emu,
                        'height_emu': height_emu,
                        'width_inches': int(width_emu) / 914400 if width_emu else None,
                        'height_inches': int(height_emu) / 914400 if height_emu else None
                    }
                    positions.append(position)
    
    # Scan main document
    scan_container(doc, 'main_document')
    
    # Scan headers and footers
    for i, section in enumerate(doc.sections):
        scan_container(section.header, 'header', i)
        scan_container(section.footer, 'footer', i)
    
    return positions

if __name__ == "__main__":
    template_path = "soc2t2-report-template-unqualified.docx"
    
    if os.path.exists(template_path):
        analyze_template_images(template_path)
        
        print("\n" + "="*50)
        print("DETAILED POSITIONS:")
        positions = get_image_positions(template_path)
        
        for i, pos in enumerate(positions):
            print(f"\nImage {i+1}:")
            print(f"  Location: {pos['container_type']}")
            if pos['section_num'] is not None:
                print(f"  Section: {pos['section_num'] + 1}")
            print(f"  Paragraph: {pos['paragraph_index'] + 1}, Run: {pos['run_index'] + 1}")
            if pos['width_inches']:
                print(f"  Size: {pos['width_inches']:.2f}\" x {pos['height_inches']:.2f}\"")
    else:
        print(f"Template file not found: {template_path}")