from docx import Document
from docx.shared import Inches
from docx.oxml.shared import qn
import os

class LogoReplacer:
    """
    Advanced logo replacement system that targets specific images by position and size
    """
    
    def __init__(self, template_path):
        self.template_path = template_path
        self.doc = Document(template_path)
        
        # Define logo contexts based on analysis
        self.logo_contexts = {
            'title_page': {
                'location': 'main_document',
                'paragraph_index': 0,  # First paragraph
                'expected_size_range': (2.0, 2.5),  # Width range in inches
                'replacement_size': Inches(2.31)
            },
            'header_small': {
                'location': 'header',
                'expected_size_range': (0.6, 0.8),  # Small header logos
                'replacement_size': Inches(0.70)
            },
            'closing_page': {
                'location': 'main_document', 
                'paragraph_index': 214,  # Last logo paragraph (0-indexed)
                'expected_size_range': (1.8, 2.0),
                'replacement_size': Inches(1.92)
            }
        }
    
    def get_image_positions(self):
        """Get all image positions in the document"""
        positions = []
        
        def scan_container(container, container_type, section_num=None):
            for para_idx, paragraph in enumerate(container.paragraphs):
                for run_idx, run in enumerate(paragraph.runs):
                    drawing_elements = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
                    
                    if drawing_elements:
                        # Get dimensions
                        width_emu = height_emu = None
                        drawings = run._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
                        
                        if drawings:
                            width_emu = drawings[0].get('cx')
                            height_emu = drawings[0].get('cy')
                        
                        position = {
                            'container_type': container_type,
                            'section_num': section_num,
                            'paragraph_index': para_idx,
                            'run_index': run_idx,
                            'width_emu': width_emu,
                            'height_emu': height_emu,
                            'width_inches': int(width_emu) / 914400 if width_emu else None,
                            'height_inches': int(height_emu) / 914400 if height_emu else None,
                            'paragraph': paragraph,
                            'run': run
                        }
                        positions.append(position)
        
        # Scan main document
        scan_container(self.doc, 'main_document')
        
        # Scan headers and footers
        for i, section in enumerate(self.doc.sections):
            scan_container(section.header, 'header', i)
            scan_container(section.footer, 'footer', i)
        
        return positions
    
    def replace_logo_by_context(self, new_logo_path, context='title_page'):
        """
        Replace logo based on predefined context (title_page, header_small, closing_page)
        """
        if context not in self.logo_contexts:
            raise ValueError(f"Unknown context: {context}. Available: {list(self.logo_contexts.keys())}")
        
        context_config = self.logo_contexts[context]
        positions = self.get_image_positions()
        replacements_made = 0
        
        for pos in positions:
            # Check if this image matches the context criteria
            matches_location = pos['container_type'] == context_config['location']
            
            # Check paragraph index if specified
            matches_paragraph = True
            if 'paragraph_index' in context_config:
                matches_paragraph = pos['paragraph_index'] == context_config['paragraph_index']
            
            # Check size range
            matches_size = False
            if pos['width_inches']:
                size_range = context_config['expected_size_range']
                matches_size = size_range[0] <= pos['width_inches'] <= size_range[1]
            
            if matches_location and matches_paragraph and matches_size:
                # Replace this image
                run = pos['run']
                run.clear()
                run.add_picture(new_logo_path, width=context_config['replacement_size'])
                replacements_made += 1
                print(f"✅ Replaced {context} logo at paragraph {pos['paragraph_index'] + 1}")
        
        return replacements_made
    
    def replace_logo_by_position(self, new_logo_path, container_type, paragraph_index, target_size_inches=None):
        """
        Replace logo by exact position
        """
        positions = self.get_image_positions()
        
        for pos in positions:
            if (pos['container_type'] == container_type and 
                pos['paragraph_index'] == paragraph_index):
                
                run = pos['run']
                run.clear()
                
                size = Inches(target_size_inches) if target_size_inches else pos['width_inches']
                run.add_picture(new_logo_path, width=size)
                
                print(f"✅ Replaced logo at {container_type} paragraph {paragraph_index + 1}")
                return True
        
        print(f"❌ No image found at {container_type} paragraph {paragraph_index + 1}")
        return False
    
    def replace_all_logos_with_size_control(self, logo_configs):
        """
        Replace multiple logos with different configurations
        
        logo_configs = {
            'title_page': '/path/to/large_logo.png',
            'header_small': '/path/to/small_logo.png', 
            'closing_page': '/path/to/medium_logo.png'
        }
        """
        total_replacements = 0
        
        for context, logo_path in logo_configs.items():
            if os.path.exists(logo_path):
                replacements = self.replace_logo_by_context(logo_path, context)
                total_replacements += replacements
            else:
                print(f"⚠️  Logo file not found: {logo_path}")
        
        return total_replacements
    
    def replace_all_with_same_logo(self, logo_path, size_mapping=None):
        """
        Replace all logos with the same image but different sizes
        
        size_mapping = {
            'title_page': 2.31,
            'header_small': 0.70,
            'closing_page': 1.92
        }
        """
        if not size_mapping:
            size_mapping = {
                'title_page': 2.31,
                'header_small': 0.70, 
                'closing_page': 1.92
            }
        
        total_replacements = 0
        
        for context in self.logo_contexts.keys():
            if context in size_mapping:
                # Temporarily update the replacement size
                original_size = self.logo_contexts[context]['replacement_size']
                self.logo_contexts[context]['replacement_size'] = Inches(size_mapping[context])
                
                replacements = self.replace_logo_by_context(logo_path, context)
                total_replacements += replacements
                
                # Restore original size
                self.logo_contexts[context]['replacement_size'] = original_size
        
        return total_replacements
    
    def save(self, output_path):
        """Save the document with replaced logos"""
        self.doc.save(output_path)
        print(f"📄 Document saved: {output_path}")

# Convenience function for SOC2 integration
def replace_soc2_logos(template_path, company_logo_path, output_path, logo_sizes=None):
    """
    Replace all SOC2 template logos with company logo
    
    Args:
        template_path: Path to SOC2 template
        company_logo_path: Path to company logo image
        output_path: Where to save the result
        logo_sizes: Dict of context -> size in inches, or None for defaults
    """
    replacer = LogoReplacer(template_path)
    
    if logo_sizes:
        replacements = replacer.replace_all_with_same_logo(company_logo_path, logo_sizes)
    else:
        # Use default sizes based on template analysis
        replacements = replacer.replace_all_with_same_logo(company_logo_path)
    
    replacer.save(output_path)
    return replacements

# Example usage
if __name__ == "__main__":
    template_path = "soc2t2-report-template-unqualified.docx"
    # company_logo = "path/to/company_logo.png"
    output_path = "soc2t2-report-with-company-logos.docx"
    
    # Example 1: Replace all with same logo, different sizes
    # replacements = replace_soc2_logos(template_path, company_logo, output_path)
    # print(f"Total replacements made: {replacements}")
    
    # Example 2: Custom size mapping
    # custom_sizes = {
    #     'title_page': 3.0,      # Larger title page logo
    #     'header_small': 0.5,    # Smaller header logos
    #     'closing_page': 2.0     # Medium closing logo
    # }
    # replacements = replace_soc2_logos(template_path, company_logo, output_path, custom_sizes)
    
    print("Advanced logo replacer ready to use!")
    print("\nAvailable logo contexts:")
    replacer = LogoReplacer(template_path)
    for context, config in replacer.logo_contexts.items():
        print(f"  - {context}: {config['replacement_size'].inches:.2f}\" (default)")